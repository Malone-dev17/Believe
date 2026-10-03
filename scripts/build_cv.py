"""Build an ATS-friendly CV (.docx + .html preview) from private/cv/master.json.

Usage:
  python scripts/build_cv.py --role "Sales Development Representative"
  python scripts/build_cv.py --all            # one CV per target role
  python scripts/build_cv.py --template       # master template, swappable sections highlighted

ATS rules followed: single column, no tables/text boxes/images, contact details in the
body (not the header), standard section names, real bullet lists, plain fonts.
"""
import argparse
import html
import json
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_COLOR_INDEX, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm

ROOT = Path(__file__).resolve().parent.parent
MASTER = ROOT / "private" / "cv" / "master.json"
OUT = ROOT / "private" / "cv" / "out"
ACCENT = RGBColor(0x1F, 0x3A, 0x68)
SWAPPABLE = ("Headline", "Profile", "Key Skills", "Key Achievements")


def slug(text):
    return re.sub(r"[^a-z0-9]+", "-", text.lower()).strip("-")


def resolve(master, role=None, variant=None):
    """Return (headline, variant dict) for a role title or variant key."""
    if role:
        variant = master["roles"].get(role)
        if not variant:
            raise SystemExit(f"Unknown role: {role}. Known: {', '.join(master['roles'])}")
    v = master["variants"][variant]
    title = role or variant.replace("_", " ").title()
    return f"{title} | {v['tagline']}", v


def sections(master, headline, v):
    """Ordered CV content as (section name, kind, payload)."""
    out = [
        ("Headline", "headline", headline),
        ("Profile", "para", v["summary"]),
        ("Key Skills", "skills", v["skills"]),
        ("Key Achievements", "bullets", v["key_achievements"]),
        ("Experience", "jobs", master["experience"]),
        ("Professional Development", "jobs", [master["development"]]),
        ("Education", "edu", master["education"]),
        ("Interests", "para", master["interests"]),
    ]
    return out


# ---------------------------------------------------------------- docx

def _bottom_border(paragraph):
    pPr = paragraph._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    b = OxmlElement("w:bottom")
    for k, val in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", "1F3A68")):
        b.set(qn(k), val)
    bdr.append(b)
    pPr.append(bdr)


def _run(p, text, bold=False, italic=False, size=None, color=None, highlight=False):
    r = p.add_run(text)
    r.bold, r.italic = bold, italic
    if size:
        r.font.size = Pt(size)
    if color:
        r.font.color.rgb = color
    if highlight:
        r.font.highlight_color = WD_COLOR_INDEX.YELLOW
    return r


def build_docx(master, content, path, template=False):
    doc = Document()
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    sec.left_margin = sec.right_margin = Cm(1.8)
    sec.top_margin = sec.bottom_margin = Cm(1.5)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(2)
    right_tab = sec.page_width - sec.left_margin - sec.right_margin

    c = master["contact"]
    p = doc.add_paragraph()
    _run(p, c["name"], bold=True, size=20, color=ACCENT)
    p = doc.add_paragraph()
    _run(p, " | ".join(c["line"]), size=10)

    for name, kind, payload in content:
        mark = template and name in SWAPPABLE
        if kind == "headline":
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            if mark:
                _run(p, "[SWAP: HEADLINE] ", bold=True, size=9, highlight=True)
            _run(p, payload, bold=True, size=12, color=ACCENT)
            continue

        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(10)
        h.paragraph_format.space_after = Pt(4)
        _run(h, name.upper(), bold=True, size=11, color=ACCENT)
        if mark:
            _run(h, f"   [SWAP: {name.upper()}]", bold=True, size=8, highlight=True)
        _bottom_border(h)

        if kind == "para":
            doc.add_paragraph(payload)
        elif kind == "skills":
            doc.add_paragraph(" | ".join(payload))
        elif kind == "bullets":
            for b in payload:
                doc.add_paragraph(b, style="List Bullet")
        elif kind == "jobs":
            for job in payload:
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(5)
                p.paragraph_format.tab_stops.add_tab_stop(right_tab, WD_TAB_ALIGNMENT.RIGHT)
                _run(p, job["title"], bold=True)
                _run(p, f" | {job['org']}")
                _run(p, f"\t{job['dates']}", italic=True)
                for b in job["bullets"]:
                    doc.add_paragraph(b, style="List Bullet")
        elif kind == "edu":
            for e in payload:
                p = doc.add_paragraph()
                p.paragraph_format.tab_stops.add_tab_stop(right_tab, WD_TAB_ALIGNMENT.RIGHT)
                _run(p, e["org"], bold=True)
                _run(p, f" | {e['detail']}")
                _run(p, f"\t{e['dates']}", italic=True)

    doc.core_properties.author = c["name"]
    doc.core_properties.title = f"{c['name']} CV"
    doc.save(path)


# ---------------------------------------------------------------- html preview

def build_html(master, content, path, template=False):
    e = html.escape
    c = master["contact"]
    parts = [f"<h1>{e(c['name'])}</h1><p class=c>{e(' | '.join(c['line']))}</p>"]
    for name, kind, payload in content:
        tag = f"<mark>SWAP: {e(name.upper())}</mark>" if template and name in SWAPPABLE else ""
        if kind == "headline":
            parts.append(f"<p class=hl>{tag} {e(payload)}</p>")
            continue
        parts.append(f"<h2>{e(name)} {tag}</h2>")
        if kind == "para":
            parts.append(f"<p>{e(payload)}</p>")
        elif kind == "skills":
            parts.append(f"<p>{e(' | '.join(payload))}</p>")
        elif kind == "bullets":
            parts.append("<ul>" + "".join(f"<li>{e(b)}</li>" for b in payload) + "</ul>")
        elif kind == "jobs":
            for j in payload:
                parts.append(f"<p class=j><b>{e(j['title'])}</b> | {e(j['org'])}<i>{e(j['dates'])}</i></p>")
                parts.append("<ul>" + "".join(f"<li>{e(b)}</li>" for b in j["bullets"]) + "</ul>")
        elif kind == "edu":
            for d in payload:
                parts.append(f"<p class=j><b>{e(d['org'])}</b> | {e(d['detail'])}<i>{e(d['dates'])}</i></p>")
    css = ("body{font:10.5pt Calibri,Arial,sans-serif;max-width:760px;margin:24px auto;padding:0 36px;color:#111;background:#fff}"
           "h1{color:#1f3a68;font-size:22pt;margin:0}.c{margin:2px 0 6px}.hl{color:#1f3a68;font-weight:700;font-size:12pt}"
           "h2{color:#1f3a68;font-size:11pt;text-transform:uppercase;border-bottom:1px solid #1f3a68;margin:14px 0 6px}"
           "ul{margin:2px 0 4px 18px;padding:0}li{margin:1px 0}.j{display:flex;gap:4px;margin:6px 0 1px}.j i{margin-left:auto}"
           "mark{font-size:8pt;font-weight:700}")
    path.write_text(f"<!doctype html><meta charset=utf-8><title>CV</title><style>{css}</style>{''.join(parts)}", encoding="utf-8")


def make(master, role=None, variant=None, template=False):
    headline, v = resolve(master, role, variant)
    content = sections(master, headline, v)
    name = master["contact"]["name"].replace(" ", "_")
    stem = "master-cv-TEMPLATE" if template else f"{name}_CV_{slug(role or variant)}"
    OUT.mkdir(parents=True, exist_ok=True)
    build_docx(master, content, OUT / f"{stem}.docx", template)
    build_html(master, content, OUT / f"{stem}.html", template)
    return OUT / f"{stem}.docx"


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    g = ap.add_mutually_exclusive_group(required=True)
    g.add_argument("--role")
    g.add_argument("--all", action="store_true")
    g.add_argument("--template", action="store_true")
    a = ap.parse_args()
    m = json.loads(MASTER.read_text(encoding="utf-8"))
    if a.template:
        print(make(m, variant="tech_sales", template=True))
    elif a.all:
        for r in m["roles"]:
            print(make(m, role=r))
    else:
        print(make(m, role=a.role))
